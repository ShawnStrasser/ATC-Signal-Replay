#!/usr/bin/env python
"""
software_validate.py — Standalone software validation script.

Replays source logs to controllers with new software, compares collected output
to the configured baseline, and generates an HTML report with divergence charts.

Usage:
    python software_validate.py                  # interactive, uses settings.json
    python software_validate.py --verbose        # extra debug output
    python software_validate.py --report-only    # skip replay, just run analysis + report
    python software_validate.py --report-only-fast
    python software_validate.py --settings custom_settings.json

Settings are loaded from settings.json (editable JSON file in the same folder).
"""

from __future__ import annotations

import argparse
import importlib.util
import json
import sys
import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import duckdb
import pandas as pd
import requests
from openpyxl import load_workbook

import signal_replay as sr
from signal_replay.ntcip import send_ntcip
from signal_replay.report import generate_report

COLLECTED_DB_FILENAME = "collected.db"
BaselineSource = Tuple[str, str]
COORD_PATTERNS_DIR = Path(__file__).resolve().parent / "coord_patterns"
SEQUENCE_MATCH_THRESHOLD = 95.0
TIMING_MATCH_THRESHOLD = 90.0
TIMING_MATCH_TOLERANCE_SECONDS = 0.5
MIN_REPORT_ISSUE_CONTEXT_MINUTES = 5.0


# ---------------------------------------------------------------------------
# Logging helpers
# ---------------------------------------------------------------------------
_VERBOSE = False


def _get_timestamp_col(df: pd.DataFrame) -> str:
    """Return the timestamp column name of a raw events DataFrame."""
    return next(
        (col for col in df.columns if col.lower() in ("timestamp", "time_stamp")),
        "timestamp",
    )


def log(msg: str, *, always: bool = True) -> None:
    """Print a message. If always=False, only prints in verbose mode."""
    if always or _VERBOSE:
        print(msg, flush=True)


def vlog(msg: str) -> None:
    """Verbose-only log."""
    log(msg, always=False)


def _refresh_coord_split_csvs(coord_dir: Path) -> None:
    module_path = coord_dir.parent / "coord_split_schedule.py"
    if not module_path.exists():
        log(f"WARNING: coord split generator not found: {module_path}")
        return

    try:
        spec = importlib.util.spec_from_file_location("coord_split_schedule_runtime", module_path)
        if spec is None or spec.loader is None:
            raise ImportError(f"Unable to load module spec for {module_path}")
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        module.convert_all_pattern_files(coord_dir)
    except Exception as exc:
        log(f"WARNING: Failed to refresh coord split CSVs from JSON files: {exc}")


def _load_coord_split_schedules(coord_dir: str) -> Dict[str, List[Dict[str, object]]]:
    """Refresh the coord split CSVs from their JSON pattern files, then read them."""
    path = Path(coord_dir)
    if not path.exists():
        return {}
    _refresh_coord_split_csvs(path)
    return sr.load_coord_split_schedules(path)


# ---------------------------------------------------------------------------
# Settings
# ---------------------------------------------------------------------------
def load_settings(path: Path) -> dict:
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def get_results_dir(software_dir: Path, settings: dict) -> Path:
    return software_dir / settings["results_dir"]


def get_version_collected_db_path(software_dir: Path, settings: dict, version: str) -> Path:
    return get_results_dir(software_dir, settings) / version / COLLECTED_DB_FILENAME


def get_version_logs_dir(software_dir: Path, settings: dict, version: str) -> Path:
    return get_results_dir(software_dir, settings) / version / "logs"


def resolve_baseline_db_path(software_dir: Path, settings: dict) -> Tuple[Optional[Path], str]:
    baseline_db_path = get_version_collected_db_path(
        software_dir,
        settings,
        settings["baseline_version"],
    )
    if baseline_db_path.exists():
        return baseline_db_path, settings["baseline_version"]
    return None, f"{settings['baseline_version']} (source logs)"


def resolve_baseline_source(
    tssu: str,
    software_dir: Path,
    settings: dict,
) -> Tuple[Optional[BaselineSource], str]:
    baseline_db_path, baseline_label = resolve_baseline_db_path(software_dir, settings)
    if baseline_db_path is not None:
        return ("db", str(baseline_db_path)), baseline_label

    baseline_log = find_log(tssu, software_dir / settings["logs_dir"])
    if baseline_log is None:
        return None, baseline_label
    return ("file", str(baseline_log)), baseline_label


# ---------------------------------------------------------------------------
# Catalog helpers
# ---------------------------------------------------------------------------
def read_catalog(catalog_path: Path) -> List[dict]:
    wb = load_workbook(catalog_path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    header = [str(h).strip().lower() for h in rows[0]]
    catalog: List[dict] = []
    for row in rows[1:]:
        if not any(row):
            continue
        r = dict(zip(header, [str(v).strip() if v is not None else "" for v in row]))
        raw_cl = r.get("cyclelength", "")
        raw_off = r.get("offset", "")
        cycle_length = int(float(raw_cl)) if raw_cl not in ("", "None") else 0
        offset = float(raw_off) if raw_off not in ("", "None") else 0.0
        catalog.append({
            "TSSU": r.get("tssu", ""),
            "Version": r.get("version", ""),
            "Type": r.get("type", "Similarity") or "Similarity",
            "CycleLength": cycle_length,
            "Offset": offset,
            "Notes": r.get("notes", ""),
            "ReplaceOnRerun": r.get("replace on rerun", ""),
        })
    return catalog


def find_log(tssu: str, logs_dir: Path) -> Optional[Path]:
    for ext in (".parquet", ".csv", ".db"):
        p = logs_dir / f"{tssu}{ext}"
        if p.exists():
            return p
    matches = sorted(logs_dir.glob(f"{tssu}*"))
    return matches[0] if matches else None


def find_database(tssu: str, databases_dir: Path) -> Optional[Path]:
    matches = sorted(databases_dir.glob(f"{tssu}*"))
    return matches[0] if matches else None


def normalize_target(target: str) -> str:
    parts = target.split(":")
    if len(parts) == 2:
        host, port_text = parts
        port = int(port_text)
        return f"{host}:{port}:{port}"
    if len(parts) == 3:
        host, udp_text, http_text = parts
        return f"{host}:{int(udp_text)}:{int(http_text)}"
    raise ValueError(f"Invalid controller target: {target}")


# ---------------------------------------------------------------------------
# Build suite
# ---------------------------------------------------------------------------
def build_suite(
    settings: dict,
    software_dir: Path,
    catalog: List[dict],
    conflict_pairs: dict,
) -> Tuple[sr.SoftwareTestSuite, dict]:
    """Build the full test suite and return it with the discovered input file map."""

    logs_dir = software_dir / settings["logs_dir"]
    databases_dir = software_dir / settings["databases_dir"]
    software_version = settings["software_version"]
    _baseline_db_path, baseline_version = resolve_baseline_db_path(software_dir, settings)

    file_map: Dict[str, dict] = {}
    for r in catalog:
        tssu = r["TSSU"]
        file_map[tssu] = {
            "log": find_log(tssu, logs_dir),
            "db": find_database(tssu, databases_dir),
        }

    similarity_scenarios: List[sr.TestScenario] = []
    conflict_scenarios: List[sr.TestScenario] = []
    skipped: List[str] = []

    for r in catalog:
        tssu = r["TSSU"]
        log_path = file_map[tssu]["log"]
        db_path = file_map[tssu]["db"]
        if not log_path:
            skipped.append(tssu)
            continue

        test_type = sr.TestType.CONFLICT if r["Type"].strip().lower() == "conflict" else sr.TestType.SIMILARITY
        cycle_length = r.get("CycleLength", 0) or 0
        offset = r.get("Offset", 0.0) or 0.0
        tod_align = cycle_length == 0

        kwargs: dict = {}
        if test_type == sr.TestType.CONFLICT:
            kwargs["replays"] = 25
            pairs_key = tssu if tssu in conflict_pairs else tssu.rstrip("_c") if tssu.endswith("_c") else tssu
            if pairs_key in conflict_pairs:
                kwargs["incompatible_pairs"] = conflict_pairs[pairs_key]

        scenario = sr.TestScenario(
            scenario_id=tssu,
            database_name=str(db_path) if db_path else f"{tssu}.bin",
            events_source=str(log_path),
            test_type=test_type,
            description=f"{r['Type']} | {r['Notes']}" if r["Notes"] else r["Type"],
            notes_column=r["Notes"],
            cycle_length=cycle_length,
            cycle_offset=offset,
            tod_align=tod_align,
            **kwargs,
        )
        if test_type == sr.TestType.CONFLICT:
            conflict_scenarios.append(scenario)
        else:
            similarity_scenarios.append(scenario)

    scenarios = similarity_scenarios + conflict_scenarios

    if skipped:
        log(f"Skipped (no log): {', '.join(skipped)}")

    comp = settings.get("comparison", {})
    configured_settle_minutes = comp.get("settle_minutes", 10.0)
    analysis_start_time = str(comp.get("analysis_start_time", "")).strip()
    analysis_end_time = str(comp.get("analysis_end_time", "")).strip()
    effective_settle_minutes = configured_settle_minutes
    if analysis_start_time and scenarios and all(s.tod_align for s in scenarios):
        effective_settle_minutes = 0.0

    suite = sr.SoftwareTestSuite(
        suite_name="Software Validation",
        software_version=software_version,
        baseline_version=baseline_version,
        scenarios=scenarios,
        batches=[],
        output_dir=str(software_dir / settings["results_dir"]),
        comparison_thresholds=sr.ComparisonThresholds(
            sequence_threshold=comp.get("sequence_threshold", 0.05),
            timing_threshold=comp.get("timing_threshold", 0.02),
            match_threshold=comp.get("match_threshold", 95.0),
        ),
        phase_call_similarity_threshold=comp.get("phase_call_similarity_threshold", comp.get("detector_similarity_threshold", 90.0)),
        analysis_settle_minutes=effective_settle_minutes,
        analysis_start_time=analysis_start_time,
        analysis_end_time=analysis_end_time,
        replay_latency_offset_seconds=settings.get("replay_latency_offset_seconds", sr.DEFAULT_REPLAY_LATENCY_OFFSET_SECONDS),
        replay_latency_offset_lookback_min=settings.get(
            "replay_latency_offset_lookback_min",
            settings.get("replay_latency_offset_update_min"),
        ),
        replay_latency_offset_min_samples=settings.get("replay_latency_offset_min_samples"),
    )

    return suite, file_map


def _replace_on_rerun_enabled(value: object) -> bool:
    return str(value).strip().lower() == "yes"


def _shared_collected_db_path(suite: sr.SoftwareTestSuite) -> Path:
    return Path(suite.output_dir) / suite.software_version / COLLECTED_DB_FILENAME


def _get_collected_device_ids(db_path: Path) -> set[str]:
    if not db_path.exists():
        return set()

    con = duckdb.connect(str(db_path), read_only=True)
    try:
        tables = {
            row[0].lower()
            for row in con.execute(
                "SELECT table_name FROM information_schema.tables WHERE table_schema = 'main'"
            ).fetchall()
        }
        device_ids: set[str] = set()
        if "events" in tables:
            device_ids.update(
                row[0]
                for row in con.execute(
                    "SELECT DISTINCT device_id FROM events WHERE device_id IS NOT NULL"
                ).fetchall()
            )
        if "conflicts" in tables:
            device_ids.update(
                row[0]
                for row in con.execute(
                    "SELECT DISTINCT device_id FROM conflicts WHERE device_id IS NOT NULL"
                ).fetchall()
            )
        return device_ids
    finally:
        con.close()


def _select_pending_replay_batch(
    suite: sr.SoftwareTestSuite,
    settings: dict,
    catalog: List[dict],
    existing_ids: set[str],
) -> Tuple[Optional[sr.TestBatch], List[sr.TestScenario], List[str], List[str]]:
    replace_flags = {
        row["TSSU"]: _replace_on_rerun_enabled(row.get("ReplaceOnRerun", ""))
        for row in catalog
    }
    pending_scenarios = [
        scenario
        for scenario in suite.scenarios
        if replace_flags.get(scenario.scenario_id, False) or scenario.scenario_id not in existing_ids
    ]
    if not pending_scenarios:
        return None, [], [], []

    normalized_targets = [normalize_target(target) for target in settings["controller_targets"]]
    selected = pending_scenarios[:len(normalized_targets)]
    batch = sr.TestBatch(
        batch_id="pending_replay",
        assignments={scenario.scenario_id: target for scenario, target in zip(selected, normalized_targets)},
        description=f"Pending replay batch: {len(selected)} scenario(s)",
    )
    replace_selected = [
        scenario.scenario_id
        for scenario in selected
        if replace_flags.get(scenario.scenario_id, False)
    ]
    skipped_existing = [
        scenario.scenario_id
        for scenario in suite.scenarios
        if scenario.scenario_id in existing_ids and not replace_flags.get(scenario.scenario_id, False)
    ]
    return batch, pending_scenarios[len(selected):], replace_selected, skipped_existing


def _clear_selected_devices_for_rerun(db_path: Path, scenario_ids: List[str]) -> None:
    if not scenario_ids or not db_path.exists():
        return

    sr.DatabaseManager(str(db_path)).clear_device_data(scenario_ids)


# ---------------------------------------------------------------------------
# Controller check
# ---------------------------------------------------------------------------
def parse_target(target: str) -> Tuple[str, int, int]:
    parts = target.split(":")
    if len(parts) == 2:
        ip, port_text = parts
        port = int(port_text)
        return ip, port, port
    if len(parts) == 3:
        ip, udp_text, http_text = parts
        return ip, int(udp_text), int(http_text)
    raise ValueError(f"Invalid target format: {target}")


def check_controller(target: str) -> bool:
    ip, udp_port, http_port = parse_target(target)
    try:
        send_ntcip((ip, udp_port), 1, 1, "Vehicle", timeout=1.0)
        send_ntcip((ip, udp_port), 1, 0, "Vehicle", timeout=1.0)
        snmp_ok = True
    except Exception:
        snmp_ok = False
    try:
        r = requests.get(f"http://{ip}:{http_port}/v1/asclog/xml/full", timeout=2.0, verify=False)
        http_ok = r.status_code == 200
    except Exception:
        http_ok = False
    return snmp_ok and http_ok


def wait_for_controllers(targets: List[str], labels: List[str]) -> None:
    """Poll controllers until all respond OK. Print status each cycle."""
    log("Polling controllers...")
    while True:
        with ThreadPoolExecutor() as ex:
            statuses = list(ex.map(check_controller, targets))
        parts = []
        for ok, lbl in zip(statuses, labels):
            parts.append(f"  OK   {lbl}" if ok else f"  FAIL {lbl}")
        print("\r\033[K" + "\n".join(parts), flush=True)
        if all(statuses):
            log("All controllers responding.")
            break
        time.sleep(3)
        # Move cursor up to overwrite
        print(f"\033[{len(parts)}A", end="", flush=True)


# ---------------------------------------------------------------------------
# Replay operations
# ---------------------------------------------------------------------------
def run_batch(suite: sr.SoftwareTestSuite, batch: sr.TestBatch) -> Path:
    runner = sr.BatchRunner(suite, debug=_VERBOSE)

    def auto_db_loader(db_name: str, target: str) -> bool:
        vlog(f"  [auto] DB {db_name} -> {target}")
        return True

    try:
        return runner.run_batch_once(batch, db_loader_callback=auto_db_loader)
    except KeyboardInterrupt:
        log("\nKeyboard interrupt received. Requesting replay shutdown...")
        runner.stop()
        raise


# ---------------------------------------------------------------------------
# Analysis (comparison runs in signal_replay.compare_validation)
# ---------------------------------------------------------------------------
def _extract_collected_events(
    suite: sr.SoftwareTestSuite,
    collected_db_path: Path,
    output_dir: Path,
) -> Dict[str, Path]:
    """
    Extract collected events from DuckDB into versioned parquet files.
    Returns a mapping of scenario_id -> parquet path.
    """
    output_dir.mkdir(parents=True, exist_ok=True)
    if not collected_db_path.exists():
        return {}

    extracted: Dict[str, Path] = {}
    vlog(f"  Extracting from {collected_db_path.name}: {len(suite.scenarios)} scenarios")
    con = duckdb.connect(str(collected_db_path), read_only=True)
    try:
        for scenario in suite.scenarios:
            df = con.execute(
                "SELECT * FROM events WHERE device_id = ? ORDER BY timestamp",
                [scenario.scenario_id],
            ).df()
            out_path = output_dir / f"{scenario.scenario_id}.parquet"
            df.to_parquet(out_path, index=False)
            extracted[scenario.scenario_id] = out_path
    finally:
        con.close()
    return extracted


def _load_collected_events_from_duckdb(db_path: Path | str, scenario_id: str) -> pd.DataFrame:
    """Load collected events for a single scenario directly from DuckDB (read-only)."""
    return sr.load_collected_events(db_path, scenario_id)


def _load_baseline_events(baseline_source: BaselineSource, scenario_id: str) -> pd.DataFrame:
    """Load baseline events from either a collected DuckDB or a source log file."""
    source_kind, source_path = baseline_source
    if source_kind == "db":
        return _load_collected_events_from_duckdb(source_path, scenario_id)
    if source_kind == "file":
        return sr.load_events(source_path)
    raise ValueError(f"Unsupported baseline source kind: {source_kind}")


def _normalize_device_csv_events(events_df: pd.DataFrame) -> pd.DataFrame:
    """Normalize raw events into a single human-readable CSV schema."""
    if events_df.empty:
        return pd.DataFrame(columns=["timestamp", "EventId", "Parameter"])

    df = events_df.copy()
    rename_map: Dict[str, str] = {}

    timestamp_col = _get_timestamp_col(df)
    if timestamp_col != "timestamp":
        rename_map[timestamp_col] = "timestamp"

    event_col = next(
        (col for col in df.columns if col.lower() in ("event_id", "eventid", "eventtypeid")),
        None,
    )
    if event_col is None:
        raise ValueError("Device CSV export requires an event_id/EventId/EventTypeID column")
    if event_col != "EventId":
        rename_map[event_col] = "EventId"

    parameter_col = next(
        (col for col in df.columns if col.lower() in ("parameter", "param")),
        None,
    )
    if parameter_col is None:
        raise ValueError("Device CSV export requires a parameter column")
    if parameter_col != "Parameter":
        rename_map[parameter_col] = "Parameter"

    if rename_map:
        df = df.rename(columns=rename_map)

    df["timestamp"] = pd.to_datetime(df["timestamp"])
    df["EventId"] = pd.to_numeric(df["EventId"], errors="raise").astype(int)
    df["Parameter"] = pd.to_numeric(df["Parameter"], errors="raise").astype(int)

    return df[["timestamp", "EventId", "Parameter"]].copy()


def _compute_export_shift(
    original: pd.DataFrame,
    collected: pd.DataFrame,
    *,
    original_ts_col: str,
    collected_ts_col: str,
    tod_align: bool,
    group_tolerance: float,
) -> pd.Timedelta:
    """Return the timestamp shift to apply to original events for export."""
    if original.empty or collected.empty:
        return pd.Timedelta(0)

    original[original_ts_col] = pd.to_datetime(original[original_ts_col])
    collected[collected_ts_col] = pd.to_datetime(collected[collected_ts_col])

    if tod_align:
        return collected[collected_ts_col].min().normalize() - original[original_ts_col].min().normalize()

    prep_a = sr.prepare_events_for_comparison(original)
    prep_b = sr.prepare_events_for_comparison(collected)
    if prep_a.empty or prep_b.empty:
        return pd.Timedelta(0)

    shift_sec = sr.find_temporal_offset(
        prep_a,
        prep_b,
        group_tolerance=group_tolerance,
    )
    original_start = prep_a["timestamp"].min()
    collected_start = prep_b["timestamp"].min()
    return (collected_start - original_start) - pd.Timedelta(seconds=shift_sec)


def _export_device_csvs(
    suite: sr.SoftwareTestSuite,
    collected_db_path: Path,
    baseline_sources: Dict[str, BaselineSource],
    group_tolerance: float = 0.0,
) -> Path:
    """Export a combined CSV per device with baseline + collected events.

    Each CSV lives in results/<sw_version>/device_events/<scenario_id>.csv
    and contains a ``DeviceId`` column (``baseline`` vs ``new``)
    to distinguish the two runs. Baseline timestamps are shifted onto the
    new run's absolute timeline using the same temporal offset logic used
    by the comparison / Gantt chart alignment.
    """
    out_dir = Path(suite.output_dir) / suite.software_version / "device_events"
    out_dir.mkdir(parents=True, exist_ok=True)

    if not collected_db_path.exists():
        log(f"Collected DuckDB not found for device CSV export: {collected_db_path}")
        return out_dir

    exported = 0
    for scenario in suite.scenarios:
        baseline_source = baseline_sources.get(scenario.scenario_id)
        if baseline_source is None:
            log(
                f"WARNING: Missing baseline source for {scenario.scenario_id}; "
                "skipping device CSV export for this scenario."
            )
            continue

        baseline = _load_baseline_events(baseline_source, scenario.scenario_id)
        collected = _load_collected_events_from_duckdb(collected_db_path, scenario.scenario_id)

        if baseline.empty:
            log(
                f"WARNING: Baseline log for {scenario.scenario_id} has no rows; "
                "skipping device CSV export for this scenario."
            )
            continue
        if collected.empty:
            log(
                f"WARNING: No collected events found in DuckDB for {scenario.scenario_id}; "
                "skipping device CSV export for this scenario."
            )
            continue

        original_ts_col = _get_timestamp_col(baseline)
        collected_ts_col = _get_timestamp_col(collected)
        baseline[original_ts_col] = pd.to_datetime(baseline[original_ts_col])
        collected[collected_ts_col] = pd.to_datetime(collected[collected_ts_col])

        original_shift = _compute_export_shift(
            baseline,
            collected,
            original_ts_col=original_ts_col,
            collected_ts_col=collected_ts_col,
            tod_align=scenario.tod_align,
            group_tolerance=group_tolerance,
        )
        if original_shift != pd.Timedelta(0):
            baseline[original_ts_col] = baseline[original_ts_col] + original_shift
        if scenario.tod_align:
            vlog(
                f"  {scenario.scenario_id}: TOD export date shift "
                f"{original_shift.total_seconds():+.2f}s"
            )
        elif original_shift != pd.Timedelta(0):
            vlog(
                f"  {scenario.scenario_id}: shifted original timestamps by "
                f"{original_shift.total_seconds():+.2f}s"
            )

        baseline = _normalize_device_csv_events(baseline)
        collected = _normalize_device_csv_events(collected)

        if baseline.empty or collected.empty:
            missing_side = "baseline" if baseline.empty else "collected"
            log(
                f"WARNING: Normalized device CSV export for {scenario.scenario_id} is missing "
                f"{missing_side} rows; skipping export."
            )
            continue

        baseline["DeviceId"] = "baseline"
        collected["DeviceId"] = "new"

        combined = pd.concat([baseline, collected], ignore_index=True)
        combined = combined[["DeviceId", "timestamp", "EventId", "Parameter"]]
        combined = combined.sort_values(["timestamp", "EventId", "Parameter", "DeviceId"]).reset_index(drop=True)
        csv_path = out_dir / f"{scenario.scenario_id}.csv"
        combined.to_csv(csv_path, index=False)
        exported += 1
        vlog(f"  {scenario.scenario_id} -> {csv_path.name} ({len(combined)} rows)")

    log(f"Exported {exported} device CSV(s) to {out_dir}")
    return out_dir


def run_analysis(
    suite: sr.SoftwareTestSuite,
    settings: dict,
    software_dir: Path,
    *,
    export_device_csvs: bool = True,
) -> List[sr.ScenarioResult]:
    """Compare collected output to the baseline with signal_replay.compare_validation."""
    software_version = suite.software_version
    comp = settings.get("comparison", {})
    group_tolerance = comp.get("group_tolerance", 0.0)
    max_workers = settings.get("analysis_workers", 4)

    plots_dir = Path(suite.output_dir) / software_version / "divergence_plots"
    plots_dir.mkdir(parents=True, exist_ok=True)
    collected_db_path = _shared_collected_db_path(suite)

    baseline_db_path, baseline_label = resolve_baseline_db_path(software_dir, settings)
    if baseline_db_path is not None:
        log(f"Baseline collected DB: {baseline_db_path}")
    else:
        log(
            f"Baseline collected DB not found for {settings['baseline_version']}; "
            "using source logs/ as fallback."
        )

    if collected_db_path.exists():
        log(f"Reading collected output directly from DuckDB: {collected_db_path}")
    else:
        log(f"Collected DuckDB not found: {collected_db_path}")

    baseline_sources: Dict[str, BaselineSource] = {}
    for scenario in suite.scenarios:
        baseline_source, _ = resolve_baseline_source(scenario.scenario_id, software_dir, settings)
        if baseline_source is None:
            log(
                f"WARNING: Missing baseline source for {scenario.scenario_id}; "
                "this scenario will be skipped during analysis."
            )
            continue
        baseline_sources[scenario.scenario_id] = baseline_source

    if export_device_csvs:
        # Export per-device CSVs (baseline + collected combined)
        log("Exporting per-device CSV files...")
        _export_device_csvs(suite, collected_db_path, baseline_sources, group_tolerance=group_tolerance)
    else:
        log("Skipping per-device CSV export.")

    analysis_start_time = comp.get("analysis_start_time")
    analysis_end_time = comp.get("analysis_end_time")
    if analysis_start_time:
        log(f"TOD manual analysis start time: {analysis_start_time}")
    if analysis_end_time:
        log(f"TOD manual analysis end time: {analysis_end_time}")

    if not collected_db_path.exists():
        for scenario in suite.scenarios:
            log(f"  No output for {scenario.scenario_id}, skipping")
        log("No scenarios to analyze.")
        return []

    scenarios = [s for s in suite.scenarios if s.scenario_id in baseline_sources]
    if not scenarios:
        log("No scenarios to analyze.")
        return []

    coord_patterns_dir = None
    if any(s.tod_align for s in scenarios) and COORD_PATTERNS_DIR.exists():
        _refresh_coord_split_csvs(COORD_PATTERNS_DIR)
        coord_patterns_dir = str(COORD_PATTERNS_DIR)

    validation_settings = sr.ValidationSettings.from_mapping(
        comp,
        coord_patterns_dir=coord_patterns_dir,
        baseline_label=baseline_label,
        candidate_label=software_version,
    )
    n_similarity = sum(1 for s in scenarios if s.test_type == sr.TestType.SIMILARITY)
    n_conflict = len(scenarios) - n_similarity
    log(
        f"\nAnalyzing {n_similarity} similarity scenario(s) with up to {max_workers} workers"
        + (f" and {n_conflict} conflict scenario(s) from saved output logs" if n_conflict else "")
        + "..."
    )
    sys.stdout.flush()

    results = sr.compare_validation(
        baseline_sources,
        ("db", str(collected_db_path)),
        scenarios,
        validation_settings,
        plots_dir=plots_dir,
        max_workers=max_workers,
    )

    for result in results:
        if result.test_type == sr.TestType.CONFLICT:
            status = "ERROR" if result.error else ("PASS" if result.passed else "FAIL")
            log(
                f"  {result.scenario_id}: {status}  "
                f"({len(result.conflicts_found)} conflict signature(s), "
                f"{result.runs_completed}/{result.total_runs} runs completed)"
            )
            continue
        status = "ERROR" if result.error else (
            "THROWN OUT" if result.thrown_out else ("PASS" if result.passed else "FAIL")
        )
        match_text = "n/a" if result.match_percentage is None else f"{result.match_percentage:.1f}%"
        plots_msg = f", {len(result.plot_paths)} charts" if result.plot_paths else ""
        diffs_msg = ""
        if result.phase_differences:
            diffs_msg = f"\n    Phase/overlap differences ({len(result.phase_differences)} phases):\n"
            diffs_msg += sr.format_phase_differences(
                result.phase_differences, label_a=baseline_label, label_b=software_version
            )
        if result.operational_differences:
            diffs_msg += f"\n    Transition/preempt/ped service differences ({len(result.operational_differences)} rows):\n"
            diffs_msg += sr.format_phase_differences(
                result.operational_differences, label_a=baseline_label, label_b=software_version
            )
        error_msg = f"\n    {result.error}" if result.error else ""
        log(f"  {result.scenario_id}: {match_text}  {status}  ({result.num_divergences} divergences{plots_msg}){diffs_msg}{error_msg}")

    results.sort(key=lambda r: (0 if r.test_type == sr.TestType.SIMILARITY else 1, r.scenario_id))
    return results


# ---------------------------------------------------------------------------
# Report generation
# ---------------------------------------------------------------------------
def build_report(
    results: List[sr.ScenarioResult],
    suite: sr.SoftwareTestSuite,
) -> Path:
    report_dir = Path(suite.output_dir) / suite.software_version
    report_dir.mkdir(parents=True, exist_ok=True)
    report_path = report_dir / "report.html"
    generate_report(results, suite, str(report_path))
    return report_path


# ---------------------------------------------------------------------------
# Refresh versioned collected logs
# ---------------------------------------------------------------------------
def archive_and_extract(suite: sr.SoftwareTestSuite, software_dir: Path, settings: dict) -> None:
    sw_ver = suite.software_version
    collected_db_path = _shared_collected_db_path(suite)
    if not collected_db_path.exists():
        log(f"No collected DuckDB found to export: {collected_db_path}")
        return
    output_dir = get_version_logs_dir(software_dir, settings, sw_ver)
    extracted_map = _extract_collected_events(suite, collected_db_path, output_dir)
    log(f"Collected logs refreshed in {output_dir} ({len(extracted_map)} scenario(s)).")


# ---------------------------------------------------------------------------
# Main flow
# ---------------------------------------------------------------------------
def main() -> None:
    global _VERBOSE

    parser = argparse.ArgumentParser(description="Software validation: replay, compare, report.")
    parser.add_argument("--settings", default="settings.json", help="Path to settings JSON file")
    parser.add_argument("--verbose", "-v", action="store_true", help="Enable verbose output")
    report_mode = parser.add_mutually_exclusive_group()
    report_mode.add_argument("--report-only", action="store_true", help="Skip replay, just run analysis + report")
    report_mode.add_argument(
        "--report-only-fast",
        action="store_true",
        help="Skip replay and rebuild the report without refreshing device CSV exports",
    )
    parser.add_argument("--archive", action="store_true", help="Refresh the versioned collected logs after report generation")
    parser.add_argument(
        "--settle-minutes",
        type=float,
        default=None,
        help="Override the initial minutes excluded from similarity analysis/reporting",
    )
    parser.add_argument(
        "--top-n",
        type=int,
        default=None,
        metavar="N",
        help="Limit the report to the first N scenarios (by scenario ID). Useful for quick local testing.",
    )
    args = parser.parse_args()

    _VERBOSE = args.verbose

    # signal_replay logs through the standard logging module and stays silent
    # until a handler is attached. Show its progress on stdout like before.
    sr.enable_console_logging(fmt="%(asctime)s %(levelname)s %(message)s", stream=sys.stdout)

    # Preserve the invoked workspace path on Windows rather than resolving a
    # mapped drive into its UNC share, which can break temp parquet access.
    software_dir = Path(__file__).parent
    settings_path = software_dir / args.settings
    if not settings_path.exists():
        print(f"ERROR: Settings file not found: {settings_path}", file=sys.stderr)
        sys.exit(1)

    settings = load_settings(settings_path)
    if args.settle_minutes is not None:
        settings.setdefault("comparison", {})["settle_minutes"] = args.settle_minutes
    log(f"signal_replay version: {sr.__version__}")
    log(f"Software dir:  {software_dir}")
    log(f"Settings:      {settings_path}")

    # --- Read catalog ---
    catalog_path = software_dir / settings["catalog_file"]
    catalog = read_catalog(catalog_path)
    log(f"Catalog:       {len(catalog)} rows from {catalog_path.name}")

    # --- Load conflict pairs ---
    conflict_pairs_path = software_dir / settings["conflict_pairs_file"]
    conflict_pairs: dict = {}
    if conflict_pairs_path.exists():
        with open(conflict_pairs_path, "r") as f:
            conflict_pairs = json.load(f)
        conflict_pairs = {k: [tuple(pair) for pair in v] for k, v in conflict_pairs.items()}
        vlog(f"Loaded conflict pairs for {len(conflict_pairs)} devices")

    # --- Build suite ---
    suite, file_map = build_suite(settings, software_dir, catalog, conflict_pairs)
    scenario_lookup = {scenario.scenario_id: scenario for scenario in suite.scenarios}

    if suite.analysis_start_time:
        if suite.scenarios and all(s.tod_align for s in suite.scenarios):
            log(
                f"TOD manual analysis start time: {suite.analysis_start_time} "
                "(replaces settle window for all scenarios)"
            )
        elif any(s.tod_align for s in suite.scenarios):
            log(f"Analysis settle window: {suite.analysis_settle_minutes} minutes")
            log(
                f"TOD manual analysis start time: {suite.analysis_start_time} "
                "(replaces settle window only for TOD-aligned scenarios)"
            )
        else:
            log(f"Analysis settle window: {suite.analysis_settle_minutes} minutes")
            log(
                f"TOD manual analysis start time: {suite.analysis_start_time} "
                "(no TOD-aligned scenarios currently use it)"
            )
    else:
        log(f"Analysis settle window: {suite.analysis_settle_minutes} minutes")
    if suite.analysis_end_time:
        log(f"TOD manual analysis end time: {suite.analysis_end_time}")

    log(f"Scenarios: {len(suite.scenarios)}  |  Controllers: {len(settings['controller_targets'])}")

    # --- File readiness ---
    ready = sum(1 for r in catalog if file_map[r["TSSU"]]["log"])
    log(f"Logs ready: {ready}/{len(catalog)}")

    if _VERBOSE:
        for r in catalog:
            tssu = r["TSSU"]
            lf = file_map[tssu]["log"]
            df = file_map[tssu]["db"]
            status = "OK" if lf else "--"
            log(f"  {status:>2s}  {tssu:>8s}  log={lf.name if lf else 'MISSING':<30s}  db={df.name if df else '--'}")

    try:
        # ======================================================================
        # REPLAY PHASE
        # ======================================================================
        if not args.report_only and not args.report_only_fast:
            existing_ids = _get_collected_device_ids(_shared_collected_db_path(suite))
            current_batch, remaining_pending, replace_selected, skipped_existing = _select_pending_replay_batch(
                suite,
                settings,
                catalog,
                existing_ids,
            )

            if skipped_existing:
                vlog(
                    "Already collected; skipping unless Replace on Rerun = yes: "
                    + ", ".join(skipped_existing)
                )

            if current_batch is None:
                log("\nNo pending replay devices in the current catalog. Proceeding to analysis.")
            else:
                shared_db_path = _shared_collected_db_path(suite)
                if replace_selected:
                    log(
                        "Replace on Rerun = yes; clearing existing data for: "
                        + ", ".join(replace_selected)
                    )
                    _clear_selected_devices_for_rerun(shared_db_path, replace_selected)

                log(f"\n{'='*70}")
                log(f"REPLAY: Next Pending Batch  ({len(current_batch.assignments)} scenarios)")
                log(f"{'='*70}")
                log("\nLoad these databases onto the controllers:")
                for sid, tgt in current_batch.assignments.items():
                    s = scenario_lookup[sid]
                    db_name = Path(s.database_name).name
                    extras = []
                    if s.test_type == sr.TestType.CONFLICT:
                        extras.append("CONFLICT" + (" pairs=configured" if s.incompatible_pairs else " (no pairs)"))
                    if s.cycle_length:
                        extras.append(f"CL={s.cycle_length}")
                    if s.cycle_offset:
                        extras.append(f"Off={s.cycle_offset}")
                    if not s.tod_align:
                        extras.append("tod_align=OFF")
                    extra_str = "  " + ", ".join(extras) if extras else ""
                    log(f"  {tgt:<22s}  <-  {db_name:<20s}{extra_str}")

                log("")
                input("Press ENTER when databases are loaded and controllers are ready...")

                log("\nChecking controllers...")
                db_lookup = {sid: Path(s.database_name).name for sid, s in scenario_lookup.items()}
                b_targets = list(current_batch.assignments.values())
                b_labels = [f"{tgt} / {db_lookup[sid]}" for sid, tgt in current_batch.assignments.items()]

                try:
                    wait_for_controllers(b_targets, b_labels)
                except KeyboardInterrupt:
                    log("\nController check interrupted. Continuing anyway...")

                log("\nRunning replay for the next pending devices...")
                try:
                    run_batch(suite, current_batch)
                except KeyboardInterrupt:
                    log("Replay interrupted by user. Exiting.")
                    sys.exit(130)

                if remaining_pending:
                    log(
                        "\nReplay batch complete. Remaining pending devices from the current catalog: "
                        + ", ".join(s.scenario_id for s in remaining_pending)
                    )
                    log("Run this script again to continue with the next pending devices.")
                    sys.exit(0)

                log("\nAll pending devices from the current catalog have been collected. Proceeding to analysis.")

        # ======================================================================
        # ANALYSIS PHASE
        # ======================================================================
        log(f"\n{'='*70}")
        log(f"ANALYSIS: Comparing {suite.software_version} output to {suite.baseline_version}")
        log(f"{'='*70}")

        if args.top_n is not None:
            suite.scenarios = suite.scenarios[:args.top_n]
            log(f"--top-n {args.top_n}: limiting analysis to {len(suite.scenarios)} scenario(s): {', '.join(s.scenario_id for s in suite.scenarios)}")

        export_device_csvs = not args.report_only_fast
        if args.report_only_fast:
            log("Report-only-fast mode: skipping device CSV refresh.")

        results = run_analysis(
            suite,
            settings,
            software_dir,
            export_device_csvs=export_device_csvs,
        )
        passed = sum(1 for r in results if r.passed)
        log(f"\nResults: {passed}/{len(results)} passed")

        # ======================================================================
        # REPORT
        # ======================================================================
        log(f"\n{'='*70}")
        log("REPORT: Generating HTML report")
        log(f"{'='*70}")

        report_path = build_report(results, suite)
        log(f"Report saved to: {report_path}")
        log(f"Total images embedded: {sum(len(r.plot_paths) for r in results)}")

        # ======================================================================
        # VERSIONED LOG EXPORT (optional refresh)
        # ======================================================================
        if args.archive:
            log(f"\n{'='*70}")
            log("EXPORT: Refreshing versioned collected logs")
            log(f"{'='*70}")
            archive_and_extract(suite, software_dir, settings)

        log("\nDone.")
    finally:
        pass


if __name__ == "__main__":
    main()
